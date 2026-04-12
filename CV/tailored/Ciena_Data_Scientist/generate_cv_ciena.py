#!/usr/bin/env python3
"""
Generate a tailored Word (.docx) CV for Ciena – Data Scientist – AI Agents.
Run:  python3 generate_cv_ciena.py
Output: CV_Ciena_Data_Scientist.docx
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml
import os

# ── Colours ─────────────────────────────────────────────────────────────────
BLUE = RGBColor(0x1F, 0x4E, 0x79)
BLACK = RGBColor(0x00, 0x00, 0x00)
GREY = RGBColor(0x4D, 0x4D, 0x4D)

doc = Document()

# ── Page margins ────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin = Cm(1.0)
    section.bottom_margin = Cm(1.0)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10)
style.font.color.rgb = BLACK
style.paragraph_format.space_before = Pt(0)
style.paragraph_format.space_after = Pt(0)

# ── Helpers ─────────────────────────────────────────────────────────────────
def heading(text):
    p = doc.add_paragraph()
    p.space_before = Pt(8)
    p.space_after = Pt(3)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = BLUE
    run.font.name = "Calibri"
    pPr = p._p.get_or_add_pPr()
    pPr.append(parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="1" w:color="1F4E79"/>'
        f'</w:pBdr>'
    ))


def body(text, space_after=2):
    p = doc.add_paragraph()
    p.space_after = Pt(space_after)
    r = p.add_run(text)
    r.font.size = Pt(9.5)
    r.font.name = "Calibri"
    return p


def bullet(text, bold_start=""):
    p = doc.add_paragraph(style="List Bullet")
    p.space_before = Pt(0)
    p.space_after = Pt(1)
    p.paragraph_format.left_indent = Inches(0.3)
    if bold_start:
        rb = p.add_run(bold_start)
        rb.bold = True
        rb.font.size = Pt(9.5)
        rb.font.name = "Calibri"
    r = p.add_run(text)
    r.font.size = Pt(9.5)
    r.font.name = "Calibri"
    return p


def italic_bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.space_before = Pt(0)
    p.space_after = Pt(1)
    p.paragraph_format.left_indent = Inches(0.3)
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9.5)
    r.font.name = "Calibri"
    return p


def exp_header(title, org, dates):
    p = doc.add_paragraph()
    p.space_before = Pt(3)
    p.space_after = Pt(1)
    p.paragraph_format.tab_stops.add_tab_stop(
        Cm(17.4), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.SPACES
    )
    r1 = p.add_run(title)
    r1.bold = True
    r1.font.size = Pt(10)
    r1.font.name = "Calibri"
    r2 = p.add_run(f"  |  {org}")
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = GREY
    r2.font.name = "Calibri"
    r3 = p.add_run(f"\t{dates}")
    r3.font.size = Pt(9.5)
    r3.font.color.rgb = GREY
    r3.font.name = "Calibri"
    r3.italic = True


def pub_entry(num, title, authors, status, bullets_list):
    p = doc.add_paragraph()
    p.space_before = Pt(3)
    p.space_after = Pt(0)
    r1 = p.add_run(f"[{num}] ")
    r1.bold = True
    r1.font.size = Pt(9.5)
    r1.font.color.rgb = BLUE
    r1.font.name = "Calibri"
    r2 = p.add_run(title)
    r2.italic = True
    r2.font.size = Pt(9.5)
    r2.font.name = "Calibri"
    pa = doc.add_paragraph()
    pa.space_after = Pt(1)
    ra = pa.add_run(authors)
    ra.italic = True
    ra.font.size = Pt(9)
    ra.font.color.rgb = GREY
    ra.font.name = "Calibri"
    rs = pa.add_run(f"  ({status})")
    rs.italic = True
    rs.font.size = Pt(9)
    rs.font.color.rgb = GREY
    rs.font.name = "Calibri"
    for b in bullets_list:
        bullet(b)


def conf_entry(header_text, detail_text):
    p = doc.add_paragraph()
    p.space_before = Pt(2)
    p.space_after = Pt(0)
    r = p.add_run(header_text)
    r.bold = True
    r.font.size = Pt(9.5)
    r.font.name = "Calibri"
    p2 = doc.add_paragraph()
    p2.space_after = Pt(1)
    p2.paragraph_format.left_indent = Inches(0.2)
    r2 = p2.add_run(detail_text)
    r2.italic = True
    r2.font.size = Pt(9)
    r2.font.name = "Calibri"


# ============================================================================
#  HEADER
# ============================================================================
tp = doc.add_paragraph()
tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
tp.space_after = Pt(0)
nr = tp.add_run("GEORGES PARFAIT DJIMEFO KAPEN")
nr.bold = True
nr.font.size = Pt(16)
nr.font.color.rgb = BLUE
nr.font.name = "Calibri"

sp = doc.add_paragraph()
sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
sp.space_after = Pt(1)
sr = sp.add_run("Data Scientist | AI Agent Developer")
sr.font.size = Pt(10)
sr.font.color.rgb = GREY
sr.font.name = "Calibri"

cp = doc.add_paragraph()
cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
cp.space_after = Pt(2)
cr = cp.add_run(
    "Montréal, Québec, Canada  ·  +1 514-245-3182  ·  gpdk2025@gmail.com"
)
cr.font.size = Pt(9.5)
cr.font.color.rgb = GREY
cr.font.name = "Calibri"

# ============================================================================
#  PROFILE
# ============================================================================
heading("Profile")
body(
    "Data scientist with expertise in AI agents, Large Language Models, "
    "advanced analytics, and machine learning applied to telecom and "
    "networking environments. PhD expected Summer 2026. Proven ability to "
    "design, deploy, and evaluate production-grade LLM-based agentic systems "
    "including RAG pipelines, orchestration frameworks, and autonomous "
    "decision engines on AWS. Strong foundation in statistical modeling, "
    "data engineering, and cross-functional collaboration.",
    space_after=3,
)

# ============================================================================
#  TECHNICAL SKILLS
# ============================================================================
heading("Technical Skills")
skills = [
    ("LLM / Agents",
     "LangChain, LangGraph, LangSmith, MCP, RAG (FAISS), CRAG, RAGAS, "
     "Claude (Haiku, Opus), Kiro, Prompt Engineering"),
    ("AI / ML",
     "PyTorch, TensorFlow, Scikit-Learn, GNN, DQN, MARL, NLP, "
     "Computer Vision, MLflow"),
    ("Data",
     "PySpark, Databricks, Snowflake, Delta Lake, Pandas, SQL, "
     "Relational Databases, ETL/ELT"),
    ("Cloud",
     "AWS (SageMaker, Bedrock, EKS, Lambda, CloudWatch, S3, ECR, "
     "CloudFormation), Kubernetes, Docker"),
    ("Languages", "Python, Java, C++, SQL, R, JavaScript, TypeScript"),
    ("DevOps", "CI/CD (GitHub Actions, CodeBuild, OIDC), Git, Jenkins, "
     "GitLab, Terraform"),
    ("Visualization",
     "Streamlit, D3.js, Plotly, Dash, Power BI, Tableau"),
    ("Methods",
     "Agile/Scrum, TDD, Statistical Modeling, MLOps, "
     "Infrastructure as Code, Design Patterns"),
]
for cat, items in skills:
    p = doc.add_paragraph()
    p.space_after = Pt(1)
    r1 = p.add_run(f"{cat}: ")
    r1.bold = True
    r1.font.size = Pt(9.5)
    r1.font.name = "Calibri"
    r2 = p.add_run(items)
    r2.font.size = Pt(9.5)
    r2.font.name = "Calibri"

# ============================================================================
#  RELEVANT EXPERIENCE
# ============================================================================
heading("Relevant Experience")

exp_header("Machine Learning Intern — AI Agents & Telecom Analytics",
           "Ericsson, Montréal", "2026")
bullet("Designed and deployed a production RAG Autonomous Agentic AI system "
       "for telecom radio unit log anomaly detection, serving Ericsson's "
       "enterprise clients.")
bullet("Trained a custom Dual-Attention SLM (522M params) achieving F1 "
       "90.6%, 99.97% accuracy, and 1 ms/log inference latency on AWS "
       "SageMaker.")
bullet("Built a 6-tool ReAct agent orchestrated by Claude Haiku 4.5 on "
       "Bedrock: SageMaker scoring, FAISS RAG (631 patterns), RF physics "
       "world model, pattern classification, CloudWatch health monitoring, "
       "and on-prem EricAI.")
bullet("Implemented self-correcting RAG (CRAG) with quality grading, RAGAS "
       "faithfulness/relevancy evaluation, and evidence-based meta-cognition "
       "confidence scoring (HIGH/MEDIUM/LOW).")
bullet("Defined and applied evaluation metrics for LLM-based systems: "
       "latency, faithfulness, relevancy, confidence calibration, and "
       "business-impact scoring.")
bullet("Engineered a production-scale map-reduce pipeline: triage → cluster "
       "by failure signature → parallel agent investigation → adaptive "
       "cross-cluster synthesis. Handles 120+ anomalies in ~12 min.")
bullet("Deployed the agent as a K8s pod on EKS exposed via Cloudflare "
       "tunnel with a Streamlit UI accessible from any device on the "
       "corporate network.")
bullet("Built an industry-standard event pipeline: S3 upload → Lambda → "
       "SQS → K8s worker → results in S3 → SNS email alerts on critical "
       "anomalies.")
bullet("Designed a fully autonomous retraining loop: drift-triggered + "
       "weekly retraining via EventBridge, with auto-promotion only when the "
       "new model's F1 exceeds the current production model.")
bullet("Implemented a 7-stage CI/CD pipeline (GitHub Actions + AWS "
       "CodeBuild with OIDC): test → OWASP security scan → CloudFormation "
       "infra → staging gate → deploy → smoke test → auto-rollback.")
italic_bullet("Stack: Python, AWS (SageMaker, Bedrock, EKS, Lambda, SQS, "
              "SNS, S3, ECR, CloudFormation, CloudWatch), LangChain, "
              "LangGraph, LangSmith, FAISS, MCP, Claude, Kiro, Docker, "
              "Kubernetes, CI/CD, Terraform.")

exp_header("AI Developer I — Data Science & Analytics",
           "Intact Financial Corporation, Montréal", "2022–2024")
bullet("Analyzed and processed large-scale datasets to uncover patterns "
       "and deliver actionable insights.")
bullet("Monitored ML pipelines and developed automated testing solutions "
       "for data quality assurance.")
bullet("Partnered with cross-functional stakeholders to translate "
       "requirements into scalable data solutions.")
italic_bullet("Stack: PySpark, SQL, Databricks, Snowflake, AWS, Shell, Git.")

exp_header("Data Engineering Intern — Advanced Analytics",
           "Intact Financial Corporation, Montréal", "2022")
bullet("Built and maintained data pipelines processing large-scale insurance "
       "datasets (PySpark, SQL, Databricks, Snowflake, AWS).")

# ============================================================================
#  ADDITIONAL EXPERIENCE
# ============================================================================
heading("Additional Experience")

additional = [
    ("Lecturer", "Polytechnique Montréal, Montréal", "2025–2026",
     ["Teach the Python procedural programming course."]),
    ("Lecturer", "Polytechnique Montréal, Montréal", "2022–2023",
     ["Teach the Python procedural programming course."]),
    ("Scientific Assistant — MOOC AI",
     "Université de Montréal, Montréal", "2021–2022",
     ["Design courses and practical work in supervised learning (Python)."]),
    ("Contractual Statistical Economist Engineer",
     "Ministry of Tourism and Recreation, Cameroon", "2012–2014",
     ["Analyze complex datasets via statistical "
      "modeling."]),
]
for title, org, dates, bullets_list in additional:
    exp_header(title, org, dates)
    for b in bullets_list:
        bullet(b)

# ============================================================================
#  SELECTED PROJECTS
# ============================================================================
heading("Selected Projects")

projects = [
    ("NLP Project", "Polytechnique Montréal, Montréal", "2021",
     ['"Extraction of keywords from texts as well as their type" (Python).']),
    ("Data Visualization Project — Chaire de recherche en fiscalité et en "
     "finances publiques", "Polytechnique Montréal, Montréal", "2021",
     ['"Develop a web application to compare taxes and provincial and '
      'federal tax rates by year."',
      "Implement a mockup with Plotly.js, D3."]),
    ("Research Initiation — Deep Learning",
     "Polytechnique Montréal, Montréal", "2020–2021",
     ['"Modeling of Voice Recognition in Android systems to secure bank '
      'payments."',
      "Model with Deep Learning (Python); implement and deploy on Android "
      "mobile (Java)."]),
    ("Programming Intern — OCR", "FLEX GROUP, Laval", "2019–2020",
     ['"Modeling of Automatic Optical Character Recognition."',
      "Implement models in Python, PHP."]),
]
for title, org, dates, bullets_list in projects:
    exp_header(title, org, dates)
    for b in bullets_list:
        bullet(b)

# ============================================================================
#  PUBLICATIONS
# ============================================================================
heading("Publications")

pub_entry(
    1,
    "Energy-Efficient and Near Real-Time Routing Algorithm for Disaster "
    "Monitoring in Wireless Sensor Networks",
    "G. P. Djimefo Kapen, R. Al Mallah, S. Pierre — Polytechnique Montréal",
    "Submitted",
    [
        "Designed GAPF, a novel meta-learning framework fusing MARL policies "
        "via a GCN encoder with only 1,473 parameters.",
        "Achieved ≥98% packet delivery, 2–3.5× less energy, 13× less memory "
        "than baselines.",
    ],
)

pub_entry(
    2,
    "Carbon-Aware Autonomous Data Reduction and Self-Healing Routing for "
    "6G-Integrated Disaster Sensor Networks",
    "G. P. Djimefo Kapen, R. Al Mallah, S. Pierre — Polytechnique Montréal",
    "Submitted",
    [
        "Proposed CALASH, the first protocol jointly optimizing Lifecycle "
        "Carbon Intensity via a hybrid Lyapunov–DQN engine and MAPE-K "
        "self-healing over 6G.",
        "Achieved 60% longer network lifetime and 33% lower LCI than "
        "baselines.",
    ],
)

pub_entry(
    3,
    "Zero-Trust Shielded Graph-Structured Decision Control for Secure "
    "Autonomous Recovery in AI-Native 6G Networks",
    "G. P. Djimefo Kapen, R. Al Mallah, S. Pierre — Polytechnique Montréal",
    "Submitted",
    [
        "Designed CASTER-ZT, the first framework combining GNN-based "
        "decision making with formally bounded zero-trust admissibility "
        "(10 proven properties).",
        "Achieved 74–98% rogue detection, zero false positives, 3 ms latency.",
    ],
)

# ============================================================================
#  EDUCATION
# ============================================================================
heading("Education")

education = [
    ("PhD in Computer Engineering (Summer 2026)",
     "Polytechnique Montréal, Montréal, Canada",
     ["AI-driven routing, carbon-aware networking, and zero-trust security "
      "for 6G sensor networks."]),
    ("Engineering Diploma in Software Engineering",
     "Polytechnique Montréal, Montréal, Canada",
     ["Concentration: Artificial Intelligence — Data Sciences."]),
    ("Master's Degree in Applied Statistics",
     "Advanced National School of Engineering (ENSPY), Yaoundé, Cameroon",
     ["Equivalency: Master's in Probability — Statistics."]),
    ("Master's Degree in Mathematics",
     "University of Yaoundé I, Yaoundé, Cameroon",
     ["Equivalency: Specialized Higher Studies Diploma (DESS) in "
      "Mathematics."]),
]
for degree, school, details in education:
    p = doc.add_paragraph()
    p.space_before = Pt(2)
    p.space_after = Pt(0)
    r = p.add_run(degree)
    r.bold = True
    r.font.size = Pt(9.5)
    r.font.name = "Calibri"
    p2 = doc.add_paragraph()
    p2.space_after = Pt(0)
    r2 = p2.add_run(school)
    r2.italic = True
    r2.font.size = Pt(9)
    r2.font.color.rgb = GREY
    r2.font.name = "Calibri"
    for d in details:
        bullet(d)

# ============================================================================
#  CERTIFICATIONS
# ============================================================================
heading("Licenses & Certifications")

certs = [
    ("AWS Certified Cloud Practitioner (CLF-C01)",
     "Certification Prep Path (Pluralsight), Montréal"),
    ("AWS Certified Developer — Associate (DVA-C02)",
     "Certification Prep Path (Pluralsight), Montréal"),
]
for cert_name, prep in certs:
    p = doc.add_paragraph(style="List Bullet")
    p.space_before = Pt(1)
    p.space_after = Pt(1)
    p.paragraph_format.left_indent = Inches(0.3)
    r1 = p.add_run(cert_name)
    r1.font.size = Pt(9.5)
    r1.font.name = "Calibri"
    r2 = p.add_run(f"\n{prep}")
    r2.italic = True
    r2.font.size = Pt(9)
    r2.font.color.rgb = GREY
    r2.font.name = "Calibri"

# ============================================================================
#  CONFERENCES
# ============================================================================
heading("Conferences & Seminars")

conf_entry(
    "Oral Presentation — PolyCongrès 2026, Polytechnique Montréal, "
    "Montréal (Mar–Apr 2026)",
    '"Energy-Efficient and Near Real-Time Routing for Disaster Monitoring '
    'in Wireless Sensor Networks — Graph-Attentive Policy Fusion (GAPF)." '
    'Georges Parfait Djimefo, Ranwa Al Mallah, Samuel Pierre.',
)
conf_entry(
    "Oral Presentation — 20th Conf. African Assoc. of Insect Scientists, "
    "Yaoundé, Cameroon (Oct 2013)",
    'Djimefo P., Tchuente M. & Le Gall P. — "Biogeography of communities '
    'of insects using clustering methods."',
)
conf_entry(
    "Poster — 29th Annual Conf. Int. Society for Clinical Biostatistics, "
    "Copenhagen, Denmark (Aug 2008)",
    '"HIV incidence estimation from repeated cross-sectional surveys: a '
    'comparative study of parametric and non-parametric approaches by '
    'simulation."',
)

# ============================================================================
#  VOLUNTEERING
# ============================================================================
heading("Volunteering")
bullet("Volunteer at the IEEE branch of Polytechnique Montréal, "
       "Montréal, 2018.")
bullet("Volunteer at the ConFoo conference, Montréal, 2018.")

# ============================================================================
#  SAVE
# ============================================================================
out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                        "CV_Ciena_Data_Scientist.docx")
doc.save(out_path)
print(f"✅  Word CV saved to: {out_path}")
